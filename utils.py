import os
from pdf2image import convert_from_path
from PIL import Image
import pytesseract
from docx import Document

# from fpdf import FPDF


def create_or_empty_dir(directory):
    """
    Create or empty the specified directory.

    Args:
            directory (str): The directory path.
    """
    if os.path.exists(directory):
        # Empty the directory if it already exists
        for filename in os.listdir(directory):
            file_path = os.path.join(directory, filename)
            os.remove(file_path)
    else:
        # Create the directory if it doesn't exist
        os.makedirs(directory)


def convert_pdf_to_images(input_pdf, output_dir):
    """
    Convert a PDF file to a series of images.

    Args:
            input_pdf (str): The path to the input PDF file.
            output_dir (str): The directory to save the converted images.
    """
    pages = convert_from_path(input_pdf)

    # Save each page as a JPEG file using Pillow
    for i, page in enumerate(pages):
        image_path = os.path.join(output_dir, f"page_{i}.jpg")
        page.save(image_path, "JPEG")


def extract_text_from_image(image_path):
    """
    Extract text from an image using OCR (Optical Character Recognition).

    Args:
            image_path (str): The path to the input image file.

    Returns:
            str: The extracted text from the image.
    """
    image = Image.open(image_path)
    text = pytesseract.image_to_string(image)
    return text


def create_docx_with_text(image_folder, output_docx):
    """
    Create a Word document (.docx) with text extracted from images.

    Args:
            image_folder (str): The directory containing the input images.
            output_docx (str): The path to save the output Word document.
    """
    document = Document()
    for filename in sorted(
        os.listdir(image_folder), key=lambda x: int(x.split("_")[1].split(".")[0])
    ):
        if filename.endswith(".png") or filename.endswith(".jpg"):
            image_path = os.path.join(image_folder, filename)
            text = extract_text_from_image(image_path)
            text = text.encode("utf-8", "ignore").decode("latin-1", "ignore")
            document.add_paragraph(text)
    document.save(output_docx)


# def create_pdf_with_text(image_folder, output_pdf):
#     """
#     Create a PDF document with text extracted from images.

#     Args:
#                     image_folder (str): The directory containing the input images.
#                     output_pdf (str): The path to save the output PDF document.
#     """

#     pdf = FPDF()
#     for filename in sorted(
#         os.listdir(image_folder), key=lambda x: int(x.split("_")[1].split(".")[0])
#     ):
#         if filename.endswith(".png") or filename.endswith(".jpg"):
#             image_path = os.path.join(image_folder, filename)
#             text = extract_text_from_image(image_path)
#             pdf.add_page()
#             pdf.set_font("Arial", size=12)
#             pdf.cell(0, 10, txt=text, ln=1)
#     pdf.output(output_pdf)
